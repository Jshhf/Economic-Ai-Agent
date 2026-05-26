from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path(r"D:\working\简历\卢家豪-AI-Agent项目升级版.docx")


def set_run_font(run, *, name: str = "Arial", east_asia: str = "Microsoft YaHei", size: int = 11, bold: bool = False, color: tuple[int, int, int] | None = None) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)
    run.font.size = Pt(size)
    run.font.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor(*color)


def add_para(doc: Document, text: str, *, size: int = 11, bold: bool = False, color: tuple[int, int, int] | None = None, after: int = 6, align: WD_ALIGN_PARAGRAPH | None = None) -> None:
    paragraph = doc.add_paragraph()
    if align is not None:
        paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(after)
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)


def add_bullet(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.line_spacing = 1.15
    run = paragraph.add_run(text)
    set_run_font(run, size=11)


def add_section_title(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(10)
    paragraph.paragraph_format.space_after = Pt(6)
    run = paragraph.add_run(text)
    set_run_font(run, size=14, bold=True, color=(0, 0, 0))


def add_role_line(doc: Document, title: str, meta: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run(title)
    set_run_font(run, size=12, bold=True)
    run2 = paragraph.add_run(f"  {meta}")
    set_run_font(run2, size=10, color=(90, 90, 90))


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(1)
section.bottom_margin = Inches(1)
section.left_margin = Inches(1)
section.right_margin = Inches(1)
section.start_type = WD_SECTION_START.NEW_PAGE

normal = doc.styles["Normal"]
normal.font.name = "Arial"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
normal.font.size = Pt(11)

add_para(doc, "卢家豪", size=24, bold=True, after=2, align=WD_ALIGN_PARAGRAPH.CENTER)
add_para(doc, "AI Agent / 后端开发 / AI-assisted engineering", size=11, color=(85, 85, 85), after=6, align=WD_ALIGN_PARAGRAPH.CENTER)
add_para(doc, "项目升级简历版 | 聚焦 Skill、RAG、MCP、可观测性", size=10, color=(85, 85, 85), after=16, align=WD_ALIGN_PARAGRAPH.CENTER)

add_section_title(doc, "项目经历")
add_role_line(
    doc,
    "Agent Runtime Platform for Economic Analysis",
    "2026-05 ~ 2026-05 | AI Agent 后端 / 全栈项目",
)
add_para(
    doc,
    "基于 FastAPI、OpenAI Agents SDK、SQLite、RAG、MCP 和 HTMX 构建 AI Agent Runtime 平台，将经济分析场景抽象为可切换 Skill 的任务式运行时系统。",
    after=4,
)
add_bullet(
    doc,
    "设计并实现 `Agent + Skill + Tool` 分层运行时，将原本固定的多 Agent 分析流程升级为 `economic_report`、`anomaly_investigation`、`policy_briefing` 三类可注册、可切换 Skill。",
)
add_bullet(
    doc,
    "构建本地知识 RAG 与 MCP 风格外部能力接入层，支持从方法论、指标定义、外部经济资料和研究背景中检索上下文，并在最终报告中输出来源引用。",
)
add_bullet(
    doc,
    "使用 SQLite 持久化 `analysis_jobs`、`artifacts`、`agent_runs`、`tool_call_logs`、`source_references`、`metrics_log`，实现 Job、Skill、RAG、MCP 全链路 trace、失败分类和运行指标记录。",
)
add_bullet(
    doc,
    "基于 Jinja2 + HTMX + Chart.js 搭建运行时可视化界面，支持技能选择、报告图表、来源查看、runtime call 分类和 metrics 展示，增强演示效果与调试能力。",
)
add_bullet(
    doc,
    "通过 AI-assisted engineering 方式完成方案设计、代码迭代、测试补齐与文档沉淀，同时主导架构拆分、接口设计、数据模型和验收标准定义。",
)

add_section_title(doc, "项目亮点表达")
add_bullet(doc, "这不是单纯的 prompt chaining，而是具备 Skill 层、工具层、RAG 层、MCP 层和 trace 层的 Agent 工程系统。")
add_bullet(doc, "重点不是“大模型会总结”，而是“我把 Agent 做成了能切换能力、能接外部来源、能做可观测性展示的运行时平台”。")
add_bullet(doc, "在面试中可以重点讲 Agent 编排、能力抽象、来源归因、失败分类和可观测性设计。")

add_section_title(doc, "可直接放入简历的项目描述")
add_para(doc, "Agent Runtime Platform for Economic Analysis  2026-05 ~ 2026-05", bold=True, after=2)
add_para(doc, "AI Agent 后端开发 | FastAPI、OpenAI Agents SDK、SQLite、RAG、MCP", color=(70, 70, 70), after=6)
add_bullet(doc, "设计并实现 AI Agent Runtime，将分析流程抽象为 `Agent + Skill + Tool` 分层架构，支持多种分析技能的服务化运行。")
add_bullet(doc, "构建本地知识 RAG 和 MCP 外部接入层，增强 Agent 对经济分析背景、指标解释和政策上下文的检索能力。")
add_bullet(doc, "使用 SQLite 持久化任务状态、运行 trace、来源引用和指标日志，提升系统可观测性与可调试性。")
add_bullet(doc, "提供 Web/API 双入口，支持上传数据、触发分析、查看报告、检查来源与运行链路。")

add_section_title(doc, "面试口径")
add_bullet(doc, "我主导了整个 Agent Runtime 的框架设计，把原来固定流程的 demo 改造成可扩展平台。")
add_bullet(doc, "我定义了 Skill、RAG、MCP、Trace 的职责边界，让系统既能扩展能力，又能调试和观测。")
add_bullet(doc, "我用 AI-assisted engineering 提高设计、编码、测试和文档效率，但核心架构、接口和数据模型是我自己收口的。")

add_section_title(doc, "关键词")
add_para(
    doc,
    "Agent Runtime, Skill Orchestration, Tool Calling, RAG, MCP Adapters, FastAPI, SQLite, Observability, Traceability, AI-assisted engineering",
    after=0,
)

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(str(OUT))
